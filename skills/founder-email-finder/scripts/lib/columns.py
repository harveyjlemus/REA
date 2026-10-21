"""Read any lead CSV and map its columns onto the fields the finder uses.

Known exports (Scrap.io, Apollo, Million Verifier reports) are recognized by
signature headers and mapped with their own table; anything else uses generic
aliases. The mapping is always shown in the preview before anything is spent,
and `--map field=Header` overrides any part of it."""
from __future__ import annotations

import csv
import io
import re
from dataclasses import asdict, dataclass, field
from pathlib import Path

from .config import Config

FIELDS = ("company", "website", "first_name", "last_name", "full_name", "title", "seniority",
          "email", "email_status", "business_email", "city", "state", "country", "linkedin_url",
          "company_linkedin", "industry", "employees", "segment", "closed", "all_emails_json",
          "catchall_flag", "mv_result", "mv_quality")


def norm_header(h: str) -> str:
    h = (h or "").replace("﻿", "").strip().lower()
    return " ".join(re.sub(r"[^a-z0-9()#]+", " ", h).split())


@dataclass
class Mapping:
    format_id: str
    format_label: str
    fields: dict[str, str] = field(default_factory=dict)   # canonical -> original header
    has_mv_columns: bool = False
    header_row: int = 0
    notes: list[str] = field(default_factory=list)

    def to_dict(self) -> dict:
        return asdict(self)

    @classmethod
    def from_dict(cls, d: dict) -> "Mapping":
        return cls(**d)


def _delimiter(sample: str) -> str:
    """Pick the delimiter from the first line only. (csv.Sniffer also guesses
    quoting, and guesses it wrong on exports with multi-line description cells.)"""
    first_line = sample.split("\n", 1)[0]
    counts = {d: first_line.count(d) for d in (",", ";", "\t", "|")}
    best = max(counts, key=counts.get)
    return best if counts[best] else ","


def _xlsx_rows(path: Path) -> list[list[str]]:
    """First sheet of an Excel file (Scrap.io exports come as .xlsx)."""
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise ValueError(f"{path.name} is an Excel file: run `python3 -m pip install openpyxl`, "
                         "or save it as CSV") from None
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = [["" if v is None else str(v) for v in r] for r in wb.worksheets[0].iter_rows(values_only=True)]
    finally:
        wb.close()
    return rows


def read_table(path: Path | str) -> tuple[list[str], list[list[str]]]:
    """(header, data rows) from a CSV or .xlsx. Finds the header row even when
    an export puts a title line above it."""
    path = Path(path)
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        rows = _xlsx_rows(path)
    else:
        raw = path.read_bytes()
        for enc in ("utf-8-sig", "cp1252", "latin-1"):
            try:
                text = raw.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        rows = list(csv.reader(io.StringIO(text, newline=""), delimiter=_delimiter(text[:8192])))
    if not rows:
        return [], []
    known = _all_known_headers()
    best, best_score = 0, -1
    for i, row in enumerate(rows[:6]):
        score = sum(1 for h in row if norm_header(h) in known)
        if score > best_score:
            best, best_score = i, score
    header = [h.replace("﻿", "").strip() for h in rows[best]]
    data = [r for r in rows[best + 1:] if any(c.strip() for c in r)]
    return header, data


_KNOWN: set[str] | None = None


def _all_known_headers() -> set[str]:
    global _KNOWN
    if _KNOWN is None:
        from .config import load_config
        cfg = load_config()
        names: set[str] = set()
        for fmt in cfg.aliases.get("formats", {}).values():
            names.update(norm_header(h) for h in fmt.get("signature", []))
            for aliases in fmt.get("map", {}).values():
                names.update(norm_header(a) for a in aliases)
        for aliases in cfg.aliases.get("generic", {}).values():
            names.update(norm_header(a) for a in aliases)
        names.add("name")
        _KNOWN = names
    return _KNOWN


def detect_mapping(header: list[str], cfg: Config) -> Mapping:
    by_norm: dict[str, str] = {}
    for h in header:
        by_norm.setdefault(norm_header(h), h)
    formats = cfg.aliases.get("formats", {})

    def has_all(sig: list[str]) -> bool:
        return all(norm_header(s) in by_norm for s in sig)

    has_mv = has_all(formats.get("millionverifier", {}).get("signature", []))
    fmt_id = "generic"
    for candidate in ("scrapio", "apollo"):
        if candidate in formats and has_all(formats[candidate]["signature"]):
            fmt_id = candidate
            break
    if fmt_id == "generic" and has_mv:
        fmt_id = "millionverifier"
    label = formats.get(fmt_id, {}).get("label", "Generic CSV")
    if has_mv and fmt_id != "millionverifier":
        label += " + Million Verifier results"
    m = Mapping(format_id=fmt_id, format_label=label, has_mv_columns=has_mv)

    def pick(aliases: list[str]) -> str | None:
        for a in aliases:
            h = by_norm.get(norm_header(a))
            if h is not None and h not in m.fields.values():
                return h
        return None

    fmt_map = formats.get(fmt_id, {}).get("map", {}) if fmt_id != "generic" else {}
    for fname, aliases in fmt_map.items():
        h = pick(aliases)
        if h:
            m.fields[fname] = h
    generic = cfg.aliases.get("generic", {})
    for fname in FIELDS:
        if fname in m.fields or fname in fmt_map or fname not in generic:
            continue
        h = pick(generic[fname])
        if h:
            m.fields[fname] = h
    if fmt_id in ("generic", "millionverifier") and "name" in by_norm and by_norm["name"] not in m.fields.values():
        if "company" in m.fields and "full_name" not in m.fields:
            m.fields["full_name"] = by_norm["name"]
            m.notes.append(f"'{by_norm['name']}' read as the person's name (a company column also exists)")
        elif "company" not in m.fields:
            m.fields["company"] = by_norm["name"]
            m.notes.append(f"'{by_norm['name']}' read as the company name; use --map full_name=... if it is a person")
    if "website" not in m.fields:
        m.notes.append("no website column found: domains will come from work emails only")
    return m


def apply_overrides(m: Mapping, header: list[str], overrides: list[str]) -> Mapping:
    """`field=Header` pairs; an empty Header removes the field."""
    by_norm = {norm_header(h): h for h in header}
    for item in overrides or []:
        if "=" not in item:
            raise ValueError(f"--map expects field=Header, got {item!r}")
        fname, _, col = item.partition("=")
        fname = fname.strip()
        if fname not in FIELDS:
            raise ValueError(f"unknown field {fname!r}; known: {', '.join(FIELDS)}")
        if not col.strip():
            m.fields.pop(fname, None)
            m.notes.append(f"{fname} unmapped by --map")
            continue
        h = by_norm.get(norm_header(col))
        if h is None:
            raise ValueError(f"--map {item!r}: no column named {col!r}")
        m.fields[fname] = h
        m.notes.append(f"{fname} <- '{h}' (from --map)")
    return m


def row_values(header: list[str], row: list[str], m: Mapping) -> tuple[dict, dict]:
    """(canonical values, original row as header->value)."""
    original = {h: (row[i] if i < len(row) else "") for i, h in enumerate(header)}
    canonical = {f: (original.get(h) or "").strip() for f, h in m.fields.items()}
    return canonical, original


def is_truthy(value: str | None) -> bool:
    return (value or "").strip().lower() in {"yes", "true", "1", "y", "closed", "permanently closed"}
